"""Build excel/nyc_taxi_kpis.xlsx from the saved query results and aggregates.

Sheets: KPI summary, month trend, weekday x hour pivot, zones, airports,
payment/tips, distance bands, cleaning log. Every number comes from
sql/results or data/processed.
"""
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
RES = ROOT / "sql" / "results"
AGG = ROOT / "data" / "processed"
OUT = ROOT / "excel" / "nyc_taxi_kpis.xlsx"
HEAD = PatternFill("solid", fgColor="1F5F99")


def r(name):
    return pd.read_csv(RES / f"{name}.csv")


def main():
    OUT.parent.mkdir(exist_ok=True)
    k = r("01_headline_kpis").iloc[0]
    kpis = pd.DataFrame([
        ("Cleaned trips", k.trips),
        ("Gross revenue ($, total_amount)", k.gross_revenue),
        ("Avg total per trip ($)", k.avg_total_per_trip),
        ("Avg fare ($)", k.avg_fare),
        ("Avg distance (mi)", k.avg_distance_mi),
        ("Median distance (mi)", k.median_distance_mi),
        ("Avg duration (min)", k.avg_duration_min),
        ("Fare per mile ($)", k.fare_per_mile),
        ("Avg trips per day", k.avg_trips_per_day),
    ], columns=["KPI", "Value"])
    heat = pd.read_csv(AGG / "weekday_hour_trips.csv")
    piv = heat.pivot_table(index=["iso_weekday", "weekday_name"], columns="hour_key", values="avg_trips").reset_index()
    piv = piv.drop(columns="iso_weekday").rename(columns={"weekday_name": "Weekday"})
    piv.columns = [str(c) for c in piv.columns]
    sheets = {
        "KPI Summary": kpis,
        "Month Trend": r("02_month_trend"),
        "Weekday x Hour": piv,
        "Weekday vs Weekend": r("10_weekday_vs_weekend"),
        "Top Pickup Zones": r("04_top_pickup_zones"),
        "Airports": r("06_airport_trips"),
        "Payment & Tips": r("07_tip_by_payment"),
        "Distance Bands": r("09_revenue_per_mile_by_distance"),
        "Daily Trips": pd.read_csv(AGG / "daily_trips.csv"),
        "Cleaning Log": pd.read_csv(AGG / "cleaning_log.csv"),
    }
    with pd.ExcelWriter(OUT, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
    wb = load_workbook(OUT)
    for ws in wb.worksheets:
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = HEAD
        for col in ws.columns:
            w = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(w + 2, 8), 45)
        ws.freeze_panes = "A2"
    ws = wb["KPI Summary"]
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        row[0].number_format = "#,##0.00"
    ws["D1"] = "Source: NYC TLC yellow taxi trip records, Jun-Aug 2025, after cleaning (see Cleaning Log)"
    ws = wb["Weekday x Hour"]
    ws.conditional_formatting.add("B2:Y8", ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="1F5F99"))
    ws = wb["Month Trend"]
    ch = BarChart(); ch.title = "Trips per day by month"
    ch.add_data(Reference(ws, min_col=4, min_row=1, max_row=4), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=2, min_row=2, max_row=4))
    ws.add_chart(ch, "I2")
    ws = wb["Daily Trips"]
    lc = LineChart(); lc.title = "Trips per day"; lc.width = 24
    lc.add_data(Reference(ws, min_col=3, min_row=1, max_row=ws.max_row), titles_from_data=True)
    lc.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ws.max_row))
    ws.add_chart(lc, "F2")
    ws = wb["Top Pickup Zones"]
    bc = BarChart(); bc.type = "bar"; bc.title = "Top pickup zones"; bc.height = 10
    bc.add_data(Reference(ws, min_col=3, min_row=1, max_row=16), titles_from_data=True)
    bc.set_categories(Reference(ws, min_col=1, min_row=2, max_row=16))
    ws.add_chart(bc, "H2")
    wb.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    main()
